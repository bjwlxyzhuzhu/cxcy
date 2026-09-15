# 抽取「新文科实践创新大赛」「职业规划大赛」指定官方文件 → JSONL（仅用户给的 6 份）。
import os, json, zipfile, re
import docx, pypdf

FILES = [
    (r"D:\Documents\C创新创业大赛\2025\新文科大赛\附件4+关于举办2025年全国大学生新文科实践创新大赛的通知.pdf", "全国大学生新文科实践创新大赛"),
    (r"D:\Documents\C创新创业大赛\2025\新文科大赛\附件1 参赛项目组别.docx", "全国大学生新文科实践创新大赛"),
    (r"D:\Documents\C创新创业大赛\2025\新文科大赛\附件2 参赛项目报告.docx", "全国大学生新文科实践创新大赛"),
    (r"D:\Documents\C创新创业大赛\2025\新文科大赛\附件3 项目评审指标体系.docx", "全国大学生新文科实践创新大赛"),
    (r"D:\Documents\C创新创业大赛\2025\职业规划大赛\附件\附件2：第三届大学生职业规划大赛就业赛道方案.docx", "全国大学生职业规划大赛"),
    (r"D:\Documents\C创新创业大赛\2025\职业规划大赛\附件\附件1：第三届大学生职业规划大赛成长赛道方案.docx", "全国大学生职业规划大赛"),
]
OUT = r"D:\Documents\BJ宝鸡教学\2026上\双创AI实验室智能体大赛\智能体平台\新职_提取.jsonl"

def clean(s):  # 去掉非法代理字符（emoji 残留等），避免 utf-8 写入崩溃
    return re.sub(r"[\ud800-\udfff]", "", s or "")

def from_docx(p):
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(x for x in parts if x and x.strip())

def from_pdf(p):
    return "\n".join((pg.extract_text() or "") for pg in pypdf.PdfReader(p).pages)

rows = []
for p, source in FILES:
    if not os.path.isfile(p):
        print("缺文件", p); continue
    ext = p.lower().rsplit(".", 1)[-1]
    try:
        t = from_docx(p) if ext == "docx" else from_pdf(p)
    except Exception as e:
        print("ERR", os.path.basename(p)[:30], str(e)[:60]); continue
    t = clean((t or "").strip())
    title = os.path.splitext(os.path.basename(p))[0]
    rows.append({"source": source, "title": title, "ext": ext, "path": p, "text": t})
    print(f"{len(t):>7}  [{source[:10]}] {title[:40]}")

with open(OUT, "w", encoding="utf-8") as f:
    for r in rows:
        f.write(json.dumps(r, ensure_ascii=False) + "\n")
print("总计", len(rows), "份 →", OUT)

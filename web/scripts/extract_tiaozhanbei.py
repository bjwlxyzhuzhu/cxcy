# 抽取「挑战杯」大挑/小挑 资料文本 → JSONL（docx/pdf/pptx 三格式）。
import os, glob, json, zipfile, re
import docx, pypdf

ROOTS = [
    (r"D:\Documents\C创新创业大赛\2025\大挑", "挑战杯·大挑（课外学术科技作品竞赛）"),
    (r"D:\Documents\C创新创业大赛\2026\挑战杯", "挑战杯·小挑（创业计划竞赛）"),
]
OUT = r"D:\Documents\BJ宝鸡教学\2026上\双创AI实验室智能体大赛\智能体平台\挑战杯_提取.jsonl"

def from_docx(p):
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(x for x in parts if x and x.strip())

def from_pdf(p):
    r = pypdf.PdfReader(p)
    return "\n".join((pg.extract_text() or "") for pg in r.pages)

def from_pptx(p):
    out = []
    with zipfile.ZipFile(p) as z:
        slides = sorted([n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)],
                        key=lambda s: int(re.search(r"\d+", s).group()))
        for n in slides:
            xml = z.read(n).decode("utf-8", "ignore")
            out += re.findall(r"<a:t>(.*?)</a:t>", xml)
    return "\n".join(out)

rows = []
for root, source in ROOTS:
    for p in glob.glob(os.path.join(root, "**", "*"), recursive=True):
        if not os.path.isfile(p):
            continue
        ext = p.lower().rsplit(".", 1)[-1] if "." in p else ""
        if ext not in ("docx", "pdf", "pptx"):
            continue
        if os.path.basename(p).startswith("~$"):
            continue
        try:
            t = from_docx(p) if ext == "docx" else (from_pdf(p) if ext == "pdf" else from_pptx(p))
        except Exception as e:
            print("ERR", os.path.basename(p)[:36], str(e)[:60]); continue
        t = (t or "").strip()
        title = os.path.splitext(os.path.basename(p))[0]
        if len(t) < 50:
            print(f"  跳过(空/短 {len(t)}) {title[:36]}"); continue
        rows.append({"source": source, "title": title, "ext": ext, "path": p, "text": t})
        print(f"{len(t):>7}  {title[:44]}")

with open(OUT, "w", encoding="utf-8") as f:
    for r in rows:
        f.write(json.dumps(r, ensure_ascii=False) + "\n")
print("总计", len(rows), "份 →", OUT)
